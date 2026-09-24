"""Shared drawing helpers for the paper's figures, drawn as editable PowerPoint shapes.

One deck per figure under ICLR/figure/, one slide per revision: the generator appends a
slide and writes its VERSION into the notes (set FIG_REPLACE=1 to overwrite the last slide
while iterating on the same revision). The slide is the printed size, so the exported PDF
(render.py) is a 1:1 drop-in for LaTeX.

Palette, shared by every figure (same meaning everywhere):
  RED    sink token / BOS, GPTQ short-document collapse
  ORANGE format (template) tokens
  GREY   user text, RTN, neutral
  BLUE   response tokens, fp16 / healthy
  GREEN  corrected objective (repair)
  PURPLE other model families (Fig. 3)
Text is never smaller than 6 pt at print size.
"""
import math
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

# saturated (S), text (T), pale fill (F)
RED_S, RED_T, RED_F = RGBColor(0xD2, 0x3F, 0x31), RGBColor(0xA9, 0x2C, 0x21), RGBColor(0xF9, 0xDB, 0xD7)
ORA_S, ORA_T, ORA_F = RGBColor(0xE8, 0x8E, 0x1E), RGBColor(0xA8, 0x62, 0x08), RGBColor(0xFB, 0xE6, 0xC5)
GRY_S, GRY_T, GRY_F = RGBColor(0xA0, 0xA0, 0xA0), RGBColor(0x6B, 0x6B, 0x6B), RGBColor(0xE8, 0xE8, 0xE8)
BLU_S, BLU_T, BLU_F = RGBColor(0x3B, 0x7D, 0xC0), RGBColor(0x24, 0x58, 0x8C), RGBColor(0xD3, 0xE3, 0xF4)
GRN_S, GRN_T, GRN_F = RGBColor(0x2F, 0x9E, 0x62), RGBColor(0x1C, 0x73, 0x45), RGBColor(0xD5, 0xEF, 0xDF)
PUR_S, PUR_T, PUR_F = RGBColor(0x8B, 0x6B, 0xB8), RGBColor(0x5E, 0x40, 0x8A), RGBColor(0xE6, 0xDD, 0xF2)
TEA_S, TEA_T = RGBColor(0x2A, 0x9D, 0x9F), RGBColor(0x1B, 0x6E, 0x70)
DARK = RGBColor(0x2B, 0x2B, 0x2B)
LIGHT = RGBColor(0xC4, 0xC4, 0xC4)
INK = RGBColor(0x26, 0x26, 0x26)
RULE = RGBColor(0x9A, 0x9A, 0x9A)
DASH = RGBColor(0x7A, 0x7A, 0x7A)
GRID = RGBColor(0xE3, 0xE3, 0xE3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = 'Arial'


class Fig:
    def __init__(self, path, w_cm, h_cm, version):
        self.path = path
        if os.path.exists(path):
            self.prs = Presentation(path)
            if os.environ.get('FIG_REPLACE') == '1' and len(self.prs.slides) > 0:
                lst = self.prs.slides._sldIdLst
                sid = list(lst)[-1]
                self.prs.part.drop_rel(sid.rId)
                lst.remove(sid)
            self.prs.slide_width, self.prs.slide_height = Cm(w_cm), Cm(h_cm)
        else:
            self.prs = Presentation()
            self.prs.slide_width, self.prs.slide_height = Cm(w_cm), Cm(h_cm)
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.slide.notes_slide.notes_text_frame.text = version
        self.S = self.slide.shapes
        self.W, self.H = w_cm, h_cm
        self.rect(0, 0, w_cm, h_cm, WHITE)

    def save(self):
        os.makedirs(os.path.dirname(self.path) or '.', exist_ok=True)
        self.prs.save(self.path)
        print('wrote', self.path, 'slides:', len(self.prs.slides))

    # ------------------------------------------------------------ primitives
    @staticmethod
    def _fmt(tf, s, size, color, bold, italic, align, anchor):
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        for i, ln in enumerate(str(s).split('\n')):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = ln
            r.font.name, r.font.size = FONT, Pt(size)
            r.font.bold, r.font.italic = bold, italic
            r.font.color.rgb = color

    def text(self, x, y, w, h, s, size=6.5, color=INK, bold=False, italic=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, rotate=0):
        tb = self.S.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
        self._fmt(tb.text_frame, s, size, color, bold, italic, align, anchor)
        if rotate:
            tb.rotation = rotate
        return tb

    def ctext(self, cx, cy, w, h, s, size=6.5, color=INK, bold=False, italic=False):
        """Text centred on (cx, cy)."""
        return self.text(cx - w / 2, cy - h / 2, w, h, s, size, color, bold, italic,
                         PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

    def vtext(self, vx, vy, vw, vh, s, size=6.5, color=INK, bold=False, align=PP_ALIGN.CENTER):
        """Text reading bottom-to-top inside the visual rectangle (vx, vy, vw, vh)."""
        tb = self.S.add_textbox(Cm(vx + vw / 2 - vh / 2), Cm(vy + vh / 2 - vw / 2), Cm(vh), Cm(vw))
        self._fmt(tb.text_frame, s, size, color, bold, False, align, MSO_ANCHOR.MIDDLE)
        tb.rotation = 270
        return tb

    def box(self, x, y, w, h, s='', size=6.5, fill=WHITE, edge=None, fg=INK, bold=False,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18, edge_w=0.5, dash=False):
        sh = self.S.add_shape(shape, Cm(x), Cm(y), Cm(w), Cm(h))
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            try:
                sh.adjustments[0] = radius
            except (IndexError, ValueError):
                pass
        if fill is None:
            sh.fill.background()
        else:
            sh.fill.solid()
            sh.fill.fore_color.rgb = fill
        if edge is None:
            sh.line.fill.background()
        else:
            sh.line.color.rgb = edge
            sh.line.width = Pt(edge_w)
            if dash:
                sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        sh.shadow.inherit = False
        self._fmt(sh.text_frame, s, size, fg, bold, False, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        return sh

    def rect(self, x, y, w, h, fill, edge=None, edge_w=0.4):
        return self.box(x, y, w, h, '', 4, fill, edge, shape=MSO_SHAPE.RECTANGLE, edge_w=edge_w)

    def dot(self, cx, cy, d, fill, edge=None, edge_w=0.6):
        return self.box(cx - d / 2, cy - d / 2, d, d, '', 4, fill, edge, shape=MSO_SHAPE.OVAL, edge_w=edge_w)

    def square(self, cx, cy, d, fill, edge=None, edge_w=0.6):
        return self.box(cx - d / 2, cy - d / 2, d, d, '', 4, fill, edge, shape=MSO_SHAPE.RECTANGLE, edge_w=edge_w)

    def diamond(self, cx, cy, d, fill, edge=None, edge_w=0.6):
        return self.box(cx - d / 2, cy - d / 2, d, d, '', 4, fill, edge, shape=MSO_SHAPE.DIAMOND, edge_w=edge_w)

    def line(self, x1, y1, x2, y2, color=RULE, w=0.75, head=False, tail=False, dash=False, dot=False):
        c = self.S.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
        c.line.color.rgb = color
        c.line.width = Pt(w)
        if dash:
            c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if dot:
            c.line.dash_style = MSO_LINE_DASH_STYLE.ROUND_DOT
        ln = c.line._get_or_add_ln()
        for tag, on in (('a:headEnd', tail), ('a:tailEnd', head)):
            if on:
                ln.append(ln.makeelement(qn(tag), {'type': 'triangle', 'w': 'sm', 'len': 'sm'}))
        return c

    def arrow(self, x1, y1, x2, y2, color=RULE, w=0.75, dash=False):
        return self.line(x1, y1, x2, y2, color, w, head=True, dash=dash)

    def polyline(self, pts, color=RULE, w=0.75, dash=False, head=False):
        segs = []
        for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
            segs.append(self.line(x1, y1, x2, y2, color, w, dash=dash))
        if head and segs:
            ln = segs[-1].line._get_or_add_ln()
            ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'sm', 'len': 'sm'}))
        return segs

    def elbow(self, x1, y1, x2, y2, color=RULE, w=0.75, head=True, via='h'):
        """Two-segment connector: horizontal-then-vertical (via='h') or vertical-then-horizontal."""
        mid = (x2, y1) if via == 'h' else (x1, y2)
        self.line(x1, y1, mid[0], mid[1], color, w)
        return self.line(mid[0], mid[1], x2, y2, color, w, head=head)

    def poly(self, pts, fill, edge=None, edge_w=0.4):
        """Closed polygon (freeform) through pts in cm."""
        fb = self.S.build_freeform(Cm(pts[0][0]), Cm(pts[0][1]), scale=1.0)
        fb.add_line_segments([(Cm(x), Cm(y)) for x, y in pts[1:]], close=True)
        sh = fb.convert_to_shape()
        if fill is None:
            sh.fill.background()
        else:
            sh.fill.solid()
            sh.fill.fore_color.rgb = fill
        if edge is None:
            sh.line.fill.background()
        else:
            sh.line.color.rgb = edge
            sh.line.width = Pt(edge_w)
        sh.shadow.inherit = False
        return sh

    def panel_label(self, x, y, letter, title='', size=7.5):
        """Bold '(a)' plus a short title, no container."""
        self.text(x, y, 0.5, 0.32, f'({letter})', size, INK, bold=True)
        if title:
            self.text(x + 0.5, y, 6.0, 0.32, title, size, INK, bold=True)

    # ------------------------------------------------------------ axes
    def frame(self, L, T, R, B, color=RULE, w=0.5, full=False):
        """Left and bottom spines (full=True: all four)."""
        if full:
            self.rect(L, T, R - L, B - T, None, color, w)
        else:
            self.line(L, T, L, B, color, w)
            self.line(L, B, R, B, color, w)

    def xticks(self, ax, y, vals, labels=None, size=6, color=GRY_T, tick=0.07, gap=0.08, fmt='%g'):
        for i, v in enumerate(vals):
            x = ax(v)
            self.line(x, y, x, y + tick, RULE, 0.5)
            lab = labels[i] if labels else (fmt % v)
            self.text(x - 0.6, y + gap, 1.2, 0.22, lab, size, color, align=PP_ALIGN.CENTER)

    def yticks(self, ay, x, vals, labels=None, size=6, color=GRY_T, tick=0.07, gap=0.08, fmt='%g', grid_to=None):
        for i, v in enumerate(vals):
            y = ay(v)
            self.line(x - tick, y, x, y, RULE, 0.5)
            if grid_to is not None:
                self.line(x, y, grid_to, y, GRID, 0.4)
            lab = labels[i] if labels else (fmt % v)
            self.text(x - gap - 1.2, y - 0.11, 1.2, 0.22, lab, size, color, align=PP_ALIGN.RIGHT)


class Axis:
    """Maps data to cm. lo/hi are data limits; a/b the cm extent (a < b for x, a > b for y)."""

    def __init__(self, lo, hi, a, b, log=False):
        self.lo, self.hi, self.a, self.b, self.log = lo, hi, a, b, log

    def __call__(self, v):
        if self.log:
            t = (math.log10(v) - math.log10(self.lo)) / (math.log10(self.hi) - math.log10(self.lo))
        else:
            t = (v - self.lo) / (self.hi - self.lo)
        return self.a + t * (self.b - self.a)


def _lerp(c0, c1, t):
    return tuple(int(round(c0[i] + t * (c1[i] - c0[i]))) for i in range(3))


HEAT_STOPS = [(0xFF, 0xFB, 0xF2), (0xFD, 0xD8, 0x9C), (0xF0, 0x6E, 0x3A), (0x8C, 0x14, 0x18)]   # cream - amber - orange-red - dark red


def heat(v, vmax, gamma=0.7):
    """Sequential cream-to-dark-red colour for v in [0, vmax], clipped."""
    t = max(0.0, min(1.0, v / vmax)) ** gamma
    n = len(HEAT_STOPS) - 1
    k = min(int(t * n), n - 1)
    return RGBColor(*_lerp(HEAT_STOPS[k], HEAT_STOPS[k + 1], t * n - k))


SUP = {-6: '10⁻⁶', -5: '10⁻⁵', -4: '10⁻⁴', -3: '10⁻³', -2: '10⁻²', -1: '10⁻¹', 0: '1', 1: '10', 2: '10²', 3: '10³', 4: '10⁴', 5: '10⁵'}
